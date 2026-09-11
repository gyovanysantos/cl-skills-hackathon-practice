"""
Tests for stage 3: the review workbook and confidence gate.
"""
import importlib.util
import json
import os

from openpyxl import load_workbook

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS_DIR = os.path.join(
    REPO_ROOT, "plugins", "cl-cim-financial-summary", "skills", "cim-financial-summary", "scripts"
)
FIXTURES_DIR = os.path.join(REPO_ROOT, "tests", "fixtures")

_spec = importlib.util.spec_from_file_location("stage3", os.path.join(SCRIPTS_DIR, "3_review_xlsx.py"))
stage3 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(stage3)


def _normalised():
    with open(os.path.join(FIXTURES_DIR, "normalised.expected.json")) as f:
        return json.load(f)


def test_every_value_has_a_row_in_all_values_sheet(tmp_path):
    normalised = _normalised()
    out = tmp_path / "review.xlsx"
    stage3.build_review_workbook(normalised, str(out))

    wb = load_workbook(out)
    ws = wb["All values"]
    # header + one row per value
    assert ws.max_row == len(normalised) + 1


def test_flagged_count_matches_review_true_count(tmp_path):
    normalised = _normalised()
    expected_flagged = sum(1 for e in normalised if e["review"])
    assert expected_flagged == 1  # the FY2023A ebitda_adjustments case

    out = tmp_path / "review.xlsx"
    flagged_count = stage3.build_review_workbook(normalised, str(out))
    assert flagged_count == expected_flagged


def test_needs_review_sheet_lists_only_flagged_rows(tmp_path):
    normalised = _normalised()
    out = tmp_path / "review.xlsx"
    stage3.build_review_workbook(normalised, str(out))

    wb = load_workbook(out)
    ws = wb["Needs review"]
    # header + 1 flagged row
    assert ws.max_row == 2
    assert ws.cell(row=2, column=2).value == "FY2023A"
    reason = ws.cell(row=2, column=6).value
    assert "reconcile" in reason or "disagreed" in reason


def test_confirmed_defaults_true_unless_flagged(tmp_path):
    normalised = _normalised()
    out = tmp_path / "review.xlsx"
    stage3.build_review_workbook(normalised, str(out))

    wb = load_workbook(out)
    ws = wb["All values"]
    header = [c.value for c in ws[1]]
    row_col = header.index("Row") + 1
    period_col = header.index("Period") + 1
    confirmed_col = header.index("Confirmed") + 1
    review_col = header.index("Review") + 1

    for r in range(2, ws.max_row + 1):
        review = ws.cell(row=r, column=review_col).value
        confirmed = ws.cell(row=r, column=confirmed_col).value
        if review:
            assert confirmed is False, f"row {r} is flagged but defaults to confirmed"
        else:
            assert confirmed is True, f"row {r} is not flagged but does not default to confirmed"


def test_flagged_rows_are_highlighted():
    from openpyxl.styles import PatternFill

    normalised = _normalised()
    import tempfile

    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "review.xlsx")
        stage3.build_review_workbook(normalised, out)
        wb = load_workbook(out)
        ws = wb["All values"]
        header = [c.value for c in ws[1]]
        review_col = header.index("Review") + 1
        for r in range(2, ws.max_row + 1):
            review = ws.cell(row=r, column=review_col).value
            fill = ws.cell(row=r, column=1).fill
            if review:
                assert fill.fgColor.rgb in ("00FEF3C7", "FFFEF3C7"), f"row {r} should be highlighted"
