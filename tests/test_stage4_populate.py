"""
Tests for stage 4: the enforced confirmation gate and deterministic
template population.
"""
import importlib.util
import json
import os

import pytest
from openpyxl import load_workbook

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS_DIR = os.path.join(
    REPO_ROOT, "plugins", "cl-cim-financial-summary", "skills", "cim-financial-summary", "scripts"
)
TEMPLATE_PATH = os.path.join(
    REPO_ROOT, "plugins", "cl-cim-financial-summary", "skills", "cim-financial-summary",
    "assets", "financial-summary-template.xlsx",
)
CONFIG_PATH = os.path.join(REPO_ROOT, "plugins", "cl-cim-financial-summary", "config", "mock-client.json")
FIXTURES_DIR = os.path.join(REPO_ROOT, "tests", "fixtures")


def _load_module(name, filename):
    spec = importlib.util.spec_from_file_location(name, os.path.join(SCRIPTS_DIR, filename))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


stage3 = _load_module("stage3", "3_review_xlsx.py")
stage4 = _load_module("stage4", "4_populate.py")


def _normalised():
    with open(os.path.join(FIXTURES_DIR, "normalised.expected.json")) as f:
        return json.load(f)


def _confirm_all(review_path):
    """Simulate the analyst: set Confirmed = TRUE on every row."""
    wb = load_workbook(review_path)
    ws = wb["All values"]
    headers = [c.value for c in ws[1]]
    confirmed_col = headers.index("Confirmed") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=confirmed_col, value=True)
    wb.save(review_path)


def test_populate_refuses_when_a_row_is_unconfirmed(tmp_path):
    normalised = _normalised()
    review_path = str(tmp_path / "review.xlsx")
    stage3.build_review_workbook(normalised, review_path)  # fresh export -- 1 row unconfirmed

    rows = stage4.load_review_rows(review_path)
    unconfirmed = stage4.find_unconfirmed(rows)
    assert len(unconfirmed) == 1
    assert unconfirmed[0]["Period"] == "FY2023A"

    output_path = tmp_path / "financial-summary.xlsx"
    exit_code = stage4.main(["4_populate.py", review_path, CONFIG_PATH, str(output_path)])
    assert exit_code != 0
    assert not output_path.exists(), "must write nothing while a row is unconfirmed"


def test_populate_succeeds_once_every_row_is_confirmed(tmp_path):
    normalised = _normalised()
    review_path = str(tmp_path / "review.xlsx")
    stage3.build_review_workbook(normalised, review_path)
    _confirm_all(review_path)

    output_path = tmp_path / "financial-summary.xlsx"
    exit_code = stage4.main(["4_populate.py", review_path, CONFIG_PATH, str(output_path)])
    assert exit_code == 0
    assert output_path.exists()


def test_populated_values_match_expected_and_only_entry_cells_changed(tmp_path):
    normalised = _normalised()
    review_path = str(tmp_path / "review.xlsx")
    stage3.build_review_workbook(normalised, review_path)
    _confirm_all(review_path)

    output_path = tmp_path / "financial-summary.xlsx"
    stage4.main(["4_populate.py", review_path, CONFIG_PATH, str(output_path)])

    populated = load_workbook(output_path)
    original = load_workbook(TEMPLATE_PATH)
    ws_pop = populated.active
    ws_orig = original.active

    # Every value_scaled from the fixture must land exactly in its entry cell,
    # except adjusted_ebitda, which has no entry cell -- the template computes
    # it live as a formula (reported EBITDA + adjustments).
    for entry in normalised:
        if entry["value_scaled"] is None or entry["template_row"] in stage4.FORMULA_DERIVED_ROWS:
            continue
        name = f"{entry['template_row']}__{entry['period']}"
        dn = populated.defined_names[name]
        sheet_name, addr = dn.attr_text.split("!")
        addr = addr.replace("$", "")
        cell = ws_pop[addr]
        assert cell.value == pytest.approx(entry["value_scaled"])

    # Formula rows must be untouched -- identical to the pristine template.
    for row in (5, 7, 9, 11, 12):
        for col in range(2, 13):
            assert ws_pop.cell(row=row, column=col).value == ws_orig.cell(row=row, column=col).value


def test_missing_periods_are_left_blank_not_invented(tmp_path):
    normalised = _normalised()
    review_path = str(tmp_path / "review.xlsx")
    stage3.build_review_workbook(normalised, review_path)
    _confirm_all(review_path)

    output_path = tmp_path / "financial-summary.xlsx"
    stage4.main(["4_populate.py", review_path, CONFIG_PATH, str(output_path)])

    populated = load_workbook(output_path)
    dn = populated.defined_names["net_revenue__FY2029E"]
    sheet_name, addr = dn.attr_text.split("!")
    addr = addr.replace("$", "")
    ws = populated[sheet_name.strip("'")]
    assert ws[addr].value is None


def test_adjusted_ebitda_rows_are_skipped_not_written(tmp_path):
    # adjusted_ebitda is extracted and reviewed (own confidence score) but has
    # no defined name in the template -- it's a live formula there. populate
    # must deliberately skip it, not error and not write over the formula.
    rows = [
        {"Template row": "adjusted_ebitda", "Period": "FY2024A", "Value (scaled)": 999.9},
        {"Template row": "net_revenue", "Period": "FY2024A", "Value (scaled)": 204.1},
    ]
    output_path = tmp_path / "out.xlsx"
    written, skipped_missing, skipped_derived = stage4.populate_template(rows, TEMPLATE_PATH, str(output_path))
    assert written == 1
    assert skipped_derived == 1

    wb = load_workbook(output_path)
    ws = wb.active
    assert isinstance(ws["F11"].value, str) and ws["F11"].value.startswith("="), (
        "adjusted EBITDA formula cell (FY2024A) must be untouched, not overwritten with 999.9"
    )


def test_populate_refuses_a_genuinely_unknown_template_row():
    fake_rows = [
        {
            "Template row": "some_row_that_does_not_exist",
            "Period": "FY2024A",
            "Value (scaled)": 999.9,
        }
    ]
    with pytest.raises(KeyError):
        stage4.populate_template(fake_rows, TEMPLATE_PATH, "/tmp/should-not-be-created.xlsx")
    assert not os.path.exists("/tmp/should-not-be-created.xlsx")
