#!/usr/bin/env python3
"""
Stage 3 -- review workbook with the confidence gate.

Turns normalised.json into a human-reviewable Excel workbook:
  - "All values" sheet: one row per (row, period) value, with source page,
    confidence, the Review flag, and a Confirmed column the analyst sets
    after checking anything flagged.
  - "Needs review" sheet: only the flagged rows, with a plain-English
    reason (which check failed), so the analyst knows exactly what to
    look at instead of re-checking everything.

Confirmed defaults to TRUE for anything not flagged, and FALSE for
anything flagged -- the analyst must deliberately tick it, in the Excel
file, before stage 4 will populate the template. This workbook, once
confirmed, is stage 4's ONLY input for values (the Value (scaled) column
here is what gets written) -- so an analyst's in-place correction to a
flagged number carries through automatically.

Dependencies: Python standard library plus openpyxl only.

Usage:
    python3 3_review_xlsx.py <normalised.json> [output.xlsx]
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1E2130")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAG_FILL = PatternFill("solid", fgColor="FEF3C7")  # amber highlight
NEEDS_REVIEW_FILL = PatternFill("solid", fgColor="FDE68A")

COLUMNS = [
    ("row", "Row"),
    ("period", "Period"),
    ("value", "Value (as extracted)"),
    ("unit", "Unit"),
    ("value_scaled", "Value (scaled)"),
    ("source_page", "Source page"),
    ("confidence", "Confidence"),
    ("review", "Review"),
    ("confirmed", "Confirmed"),
]

ROW_LABELS = {
    "net_revenue": "Net revenue",
    "gross_profit": "Gross profit",
    "reported_ebitda": "Reported EBITDA",
    "ebitda_adjustments": "Adjustments to EBITDA",
    "adjusted_ebitda": "Adjusted EBITDA",
}


def failure_reason(entry):
    """Plain-English reason a row is flagged, from its checks object."""
    checks = entry["checks"]
    reasons = []
    if not checks.get("schema", True):
        reasons.append("value did not parse to the declared type/unit")
    if not checks.get("arithmetic", True):
        reasons.append("does not reconcile with related lines (e.g. the EBITDA bridge)")
    if not checks.get("two_pass", True):
        reasons.append("two independent reads of the source disagreed")
    if not reasons and entry["confidence"] < 0.90:
        reasons.append(f"confidence {entry['confidence']:.0%} is below the review threshold")
    return "; ".join(reasons) if reasons else "flagged for review"


def build_review_workbook(normalised_data, output_path):
    wb = Workbook()
    ws_all = wb.active
    ws_all.title = "All values"

    for col_idx, (_, header) in enumerate(COLUMNS, start=1):
        c = ws_all.cell(row=1, column=col_idx, value=header)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL

    row_idx = 2
    flagged_entries = []
    for entry in normalised_data:
        review = bool(entry["review"])
        confirmed_default = not review  # auto-confirmed unless flagged

        values = {
            "row": ROW_LABELS.get(entry["row"], entry["row"]),
            "period": entry["period"],
            "value": entry["value"],
            "unit": entry["unit"],
            "value_scaled": entry["value_scaled"],
            "source_page": entry["source"]["page"],
            "confidence": entry["confidence"],
            "review": review,
            "confirmed": confirmed_default,
        }

        for col_idx, (key, _) in enumerate(COLUMNS, start=1):
            cell = ws_all.cell(row=row_idx, column=col_idx, value=values[key])
            if key == "confidence":
                cell.number_format = "0%"
            if key in ("value", "value_scaled"):
                cell.number_format = '"$"0.0'
            if review:
                cell.fill = FLAG_FILL

        if review:
            flagged_entries.append(entry)

        row_idx += 1

    for col_idx, (key, _) in enumerate(COLUMNS, start=1):
        width = 14 if key not in ("row",) else 22
        ws_all.column_dimensions[get_column_letter(col_idx)].width = width
    ws_all.freeze_panes = "A2"

    # "Needs review" sheet
    ws_flagged = wb.create_sheet("Needs review")
    flagged_headers = ["Row", "Period", "Value (as extracted)", "Confidence", "Source page", "Reason"]
    for col_idx, header in enumerate(flagged_headers, start=1):
        c = ws_flagged.cell(row=1, column=col_idx, value=header)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL

    if not flagged_entries:
        ws_flagged.cell(row=2, column=1, value="Nothing flagged -- every value met the confidence threshold.")
    else:
        for i, entry in enumerate(flagged_entries, start=2):
            ws_flagged.cell(row=i, column=1, value=ROW_LABELS.get(entry["row"], entry["row"]))
            ws_flagged.cell(row=i, column=2, value=entry["period"])
            v = ws_flagged.cell(row=i, column=3, value=entry["value"])
            v.number_format = '"$"0.0'
            conf = ws_flagged.cell(row=i, column=4, value=entry["confidence"])
            conf.number_format = "0%"
            ws_flagged.cell(row=i, column=5, value=entry["source"]["page"])
            ws_flagged.cell(row=i, column=6, value=failure_reason(entry))
            for col in range(1, 7):
                ws_flagged.cell(row=i, column=col).fill = NEEDS_REVIEW_FILL

    for col_idx, header in enumerate(flagged_headers, start=1):
        ws_flagged.column_dimensions[get_column_letter(col_idx)].width = 20 if header != "Reason" else 55

    wb.save(output_path)
    return len(flagged_entries)


def main(argv):
    if len(argv) < 2:
        print("usage: 3_review_xlsx.py <normalised.json> [output.xlsx]", file=sys.stderr)
        return 2

    normalised_path = argv[1]
    output_path = argv[2] if len(argv) > 2 else "review.xlsx"

    with open(normalised_path) as f:
        normalised_data = json.load(f)

    flagged_count = build_review_workbook(normalised_data, output_path)

    print(f"wrote {output_path} ({len(normalised_data)} values, {flagged_count} flagged for review)", file=sys.stderr)
    if flagged_count:
        print(
            f"Action needed: {flagged_count} value(s) need review before stage 4 can populate the "
            f"template. Open {output_path}, check the 'Needs review' sheet, correct any value in "
            f"'All values' if needed, then set Confirmed = TRUE for each flagged row.",
            file=sys.stderr,
        )

    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
