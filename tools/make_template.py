#!/usr/bin/env python3
"""
Builds the Financial Summary Excel template asset used by stages 3 and 4.

The template is code-generated (not hand-crafted) so its structure is
reproducible and reviewable in a diff. Run this whenever the template
layout changes:

    python3 tools/make_template.py

Writes:
    plugins/cl-cim-financial-summary/skills/cim-financial-summary/assets/financial-summary-template.xlsx

Layout (Sheet "Financial Summary"):
    Row 3:  period headers (US$ millions | FY2020A | ... | FY2030E)
    Row 4:  Net revenue                 -- ENTRY cell per period (populate.py writes here)
    Row 5:  Revenue growth %            -- FORMULA (period-over-period)
    Row 6:  Gross profit                -- ENTRY cell per period
    Row 7:  Gross margin %              -- FORMULA (= gross profit / net revenue)
    Row 8:  Reported EBITDA             -- ENTRY cell per period
    Row 9:  EBITDA margin %             -- FORMULA (= reported EBITDA / net revenue)
    Row 10: Adjustments to EBITDA       -- ENTRY cell per period
    Row 11: Adjusted EBITDA             -- FORMULA (= reported EBITDA + adjustments) -- the bridge
    Row 12: Adjusted EBITDA margin %    -- FORMULA (= adjusted EBITDA / net revenue)

Only rows 4, 6, 8, 10 are entry cells (one per period column), each given a
workbook-level defined name "<row_key>__<period>" so stage 4 can locate and
write to them by name without ever touching a formula cell. Adjusted EBITDA
is deliberately a FORMULA, not an entry cell: it is independently extracted
from the CIM too (for the stage-2 arithmetic consistency check -- does the
seller's own bridge add up?), but the template always computes it live so it
stays correct if an analyst edits an upstream assumption later.
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(
    HERE,
    "..",
    "plugins",
    "cl-cim-financial-summary",
    "skills",
    "cim-financial-summary",
    "assets",
    "financial-summary-template.xlsx",
)

PERIODS = [
    "FY2020A", "FY2021A", "FY2022A", "FY2023A", "FY2024A",
    "FY2025E", "FY2026E", "FY2027E", "FY2028E", "FY2029E", "FY2030E",
]

MONEY_FORMAT = '"$"0.0;("$"0.0)'
PERCENT_FORMAT = "0.0%"

HEADER_FILL = PatternFill("solid", fgColor="1E2130")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LABEL_FONT = Font(bold=True)
FORMULA_LABEL_FONT = Font(italic=True, color="6B7280")

# row -> (excel row number, label, kind) where kind is "entry" or "formula"
ROW_PLAN = [
    ("net_revenue", 4, "Net revenue", "entry"),
    ("revenue_growth_pct", 5, "Revenue growth %", "formula"),
    ("gross_profit", 6, "Gross profit", "entry"),
    ("gross_margin_pct", 7, "Gross margin %", "formula"),
    ("reported_ebitda", 8, "Reported EBITDA", "entry"),
    ("ebitda_margin_pct", 9, "EBITDA margin %", "formula"),
    ("ebitda_adjustments", 10, "Adjustments to EBITDA", "entry"),
    ("adjusted_ebitda", 11, "Adjusted EBITDA", "formula"),
    ("adjusted_ebitda_margin_pct", 12, "Adjusted EBITDA margin %", "formula"),
]

FIRST_PERIOD_COL = 2  # column B
HEADER_ROW = 3


def col_letter(period_index):
    return get_column_letter(FIRST_PERIOD_COL + period_index)


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Summary"

    ws["A1"] = "Financial Summary — Northbridge Packaging Holdings (fabricated, hackathon practice)"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + len(PERIODS))

    # Header row
    ws.cell(row=HEADER_ROW, column=1, value="US$ millions").font = HEADER_FONT
    ws.cell(row=HEADER_ROW, column=1).fill = HEADER_FILL
    for i, period in enumerate(PERIODS):
        c = ws.cell(row=HEADER_ROW, column=FIRST_PERIOD_COL + i, value=period)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="right")

    for row_key, excel_row, label, kind in ROW_PLAN:
        label_cell = ws.cell(row=excel_row, column=1, value=label)
        label_cell.font = LABEL_FONT if kind == "entry" else FORMULA_LABEL_FONT

        for i, period in enumerate(PERIODS):
            col = FIRST_PERIOD_COL + i
            letter = col_letter(i)
            cell = ws.cell(row=excel_row, column=col)

            if kind == "entry":
                cell.number_format = MONEY_FORMAT
                # Leave blank -- populate.py fills this. Give it a defined name so
                # populate.py can find it without hardcoding row/column numbers.
                name = f"{row_key}__{period}"
                wb.defined_names[name] = DefinedName(name, attr_text=f"'{ws.title}'!${letter}${excel_row}")
            else:
                cell.number_format = PERCENT_FORMAT if "pct" in row_key else MONEY_FORMAT
                if row_key.endswith("_pct"):
                    base_row = {"revenue_growth_pct": 4, "gross_margin_pct": 6, "ebitda_margin_pct": 8, "adjusted_ebitda_margin_pct": 11}[row_key]
                    if row_key == "revenue_growth_pct":
                        if i == 0:
                            continue  # no prior period to grow from
                        prev_letter = col_letter(i - 1)
                        cell.value = f"=({letter}4-{prev_letter}4)/{prev_letter}4"
                    else:
                        cell.value = f"={letter}{base_row}/{letter}4"
                elif row_key == "adjusted_ebitda":
                    cell.value = f"={letter}8+{letter}10"

    # Column widths
    ws.column_dimensions["A"].width = 28
    for i in range(len(PERIODS)):
        ws.column_dimensions[col_letter(i)].width = 11

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")


if __name__ == "__main__":
    build()
